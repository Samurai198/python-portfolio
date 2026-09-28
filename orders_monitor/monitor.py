"""Монитор новых заказов на фриланс-биржах (сейчас — FL.ru).

Читает официальные RSS-ленты проектов, оставляет только заказы
с нужными словами (python, парсер, бот...) и:
  - дописывает их в orders.xlsx,
  - по желанию присылает новые заказы в Telegram.

Уже виденные заказы запоминаются в seen.json, поэтому при повторном
запуске показываются только новые.

Запуск:
    python monitor.py                      # один раз проверить
    python monitor.py --watch 10           # проверять каждые 10 минут
    python monitor.py --telegram           # + слать новые заказы в Telegram

Для Telegram положите рядом файл .env с двумя строками:
    BOT_TOKEN=токен от @BotFather
    ADMIN_ID=ваш числовой id от @userinfobot
"""

import argparse
import html
import json
import os
import re
import time
import xml.etree.ElementTree as ET
from email.utils import parsedate_to_datetime
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook

# Биржа -> адрес RSS-ленты. Чтобы следить за ещё одной биржей, добавьте строку.
FEEDS = {
    "FL.ru": "https://www.fl.ru/rss/all.xml",
}
KEYWORDS = ["python", "питон", "парсер", "парсинг", "бот", "telegram",
            "телеграм", "excel", "скрипт", "автоматизац"]
SEEN_FILE = Path("seen.json")
XLSX_FILE = Path("orders.xlsx")
COLUMNS = ["date", "title", "budget", "category", "link", "source"]
ENV_FILE = Path(__file__).with_name(".env")


def parse_feed(xml_text, source=""):
    """Превращает RSS в список заказов."""
    root = ET.fromstring(xml_text)
    orders = []
    for item in root.iter("item"):
        title = html.unescape(item.findtext("title", "").strip())
        description = html.unescape(item.findtext("description", ""))
        description = re.sub(r"<[^>]+>", " ", description).strip()
        pub_date = item.findtext("pubDate")
        orders.append({
            "date": parsedate_to_datetime(pub_date).strftime("%d.%m.%Y %H:%M") if pub_date else "",
            "title": title,
            "budget": extract_budget(title + " " + description),
            "category": item.findtext("category", "").strip(),
            "link": item.findtext("link", "").strip(),
            "description": description,
            "source": source,
        })
    return orders


def extract_budget(text):
    """Ищет в тексте бюджет вида «Бюджет: 3 000 ₽» и возвращает число."""
    match = re.search(r"Бюджет:\s*([\d\s]+)", text)
    if not match:
        return None
    digits = re.sub(r"\D", "", match.group(1))
    return int(digits) if digits else None


def matches(order, keywords):
    """Ищет слова только с начала слова: «бот» найдёт «боты», но не «работа»."""
    text = f"{order['title']} {order['description']} {order['category']}".lower()
    return any(re.search(r"\b" + re.escape(word), text) for word in keywords)


def load_seen():
    if SEEN_FILE.exists():
        return set(json.loads(SEEN_FILE.read_text(encoding="utf-8")))
    return set()


def save_seen(seen):
    SEEN_FILE.write_text(json.dumps(sorted(seen), ensure_ascii=False), encoding="utf-8")


def append_xlsx(orders):
    if XLSX_FILE.exists():
        wb = load_workbook(XLSX_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Заказы"
        ws.append(COLUMNS)
        for column, width in zip("ABCDEF", (17, 70, 10, 35, 50, 15)):
            ws.column_dimensions[column].width = width
    if ws["F1"].value is None:  # файл от старой версии без колонки source
        ws["F1"] = "source"
    for order in orders:
        ws.append([order[c] for c in COLUMNS])
    wb.save(XLSX_FILE)


class TelegramError(Exception):
    pass


def send_one(order, attempts=3):
    """Шлёт один заказ; при обрыве связи пробует ещё раз. Токен в ошибках скрыт."""
    token, chat_id = os.environ["BOT_TOKEN"], os.environ["ADMIN_ID"]
    budget = f"{order['budget']} ₽" if order["budget"] else "не указан"
    text = f"🆕 [{order['source']}] {order['title']}\nБюджет: {budget}\n{order['link']}"
    for attempt in range(1, attempts + 1):
        try:
            response = requests.post(f"https://api.telegram.org/bot{token}/sendMessage",
                                     data={"chat_id": chat_id, "text": text}, timeout=15)
            response.raise_for_status()
            return
        except requests.RequestException as e:
            if attempt == attempts:
                raise TelegramError(str(e).replace(token, "***")) from None
            time.sleep(5)


def send_telegram(orders):
    for order in orders:
        send_one(order)


def load_env():
    """Читает BOT_TOKEN и ADMIN_ID из файла .env, если он есть."""
    if not ENV_FILE.exists():
        return
    for line in ENV_FILE.read_text(encoding="utf-8-sig").splitlines():  # -sig: Блокнот добавляет BOM
        key, sep, value = line.partition("=")
        if sep and not line.lstrip().startswith("#"):
            os.environ.setdefault(key.strip(), value.strip())


def fetch_orders():
    """Собирает заказы со всех лент; если одна биржа недоступна, остальные работают."""
    orders = []
    for source, url in FEEDS.items():
        try:
            response = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=20)
            response.raise_for_status()
            orders.extend(parse_feed(response.content, source))
        except (requests.RequestException, ET.ParseError) as e:
            print(f"{source}: не удалось получить ленту ({e})")
    return orders


def check(keywords, telegram):
    seen = load_seen()
    new, links = [], set()
    for order in fetch_orders():
        if order["link"] not in seen and order["link"] not in links and matches(order, keywords):
            new.append(order)
            links.add(order["link"])
    delivered = []
    for order in new:
        if telegram:
            try:
                send_one(order)
            except TelegramError as e:
                print(f"Не удалось отправить в Telegram: {e}\nОстальные заказы отправлю при следующей проверке.")
                break
        delivered.append(order)
    # В «виденные» попадают только доставленные заказы, остальные придут в следующий раз
    if delivered:
        append_xlsx(delivered)
        seen.update(o["link"] for o in delivered)
        save_seen(seen)
    for order in delivered:
        print(f"[{order['date']}] [{order['source']}] {order['title']}\n    {order['link']}")
    print(f"Новых подходящих заказов: {len(delivered)}")


def main():
    ap = argparse.ArgumentParser(description="Монитор заказов на фриланс-биржах")
    ap.add_argument("--watch", type=int, metavar="MIN", help="проверять каждые MIN минут")
    ap.add_argument("--telegram", action="store_true", help="слать новые заказы в Telegram")
    ap.add_argument("--keywords", help="свои слова через запятую, например: django,api")
    args = ap.parse_args()

    keywords = [w.strip().lower() for w in args.keywords.split(",")] if args.keywords else KEYWORDS
    if args.telegram:
        load_env()
        if not os.environ.get("BOT_TOKEN") or not os.environ.get("ADMIN_ID", "").isdigit():
            raise SystemExit("Для --telegram нужен файл .env с BOT_TOKEN и ADMIN_ID, см. README")
    while True:
        try:
            check(keywords, args.telegram)
        except requests.RequestException as e:
            print(f"Ошибка сети: {str(e).replace(os.environ.get('BOT_TOKEN') or '<нет>', '***')}")
        if not args.watch:
            break
        time.sleep(args.watch * 60)


if __name__ == "__main__":
    main()
