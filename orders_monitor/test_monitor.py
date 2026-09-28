from monitor import extract_budget, matches, parse_feed

SAMPLE = """<?xml version="1.0" encoding="utf-8"?>
<rss version="2.0"><channel>
  <item>
    <title><![CDATA[Написать парсер цен конкурентов (Бюджет: 3 000 ₽)]]></title>
    <link>https://www.fl.ru/projects/1/parser.html</link>
    <description><![CDATA[Нужно собрать цены с <b>двух сайтов</b> в Excel]]></description>
    <category>Программирование / Парсинг данных</category>
    <pubDate>Mon, 28 Sep 2026 10:15:00 +0300</pubDate>
  </item>
  <item>
    <title>Нарисовать логотип</title>
    <link>https://www.fl.ru/projects/2/logo.html</link>
    <description>Логотип для кофейни</description>
    <category>Дизайн</category>
  </item>
</channel></rss>"""


def test_parse_feed():
    orders = parse_feed(SAMPLE)
    assert len(orders) == 2
    first = orders[0]
    assert first["title"] == "Написать парсер цен конкурентов (Бюджет: 3 000 ₽)"
    assert first["budget"] == 3000
    assert first["date"] == "28.09.2026 10:15"
    assert "двух сайтов" in first["description"] and "<b>" not in first["description"]
    assert orders[1]["budget"] is None and orders[1]["date"] == ""


def test_filter_by_keywords():
    parser_order, logo_order = parse_feed(SAMPLE)
    assert matches(parser_order, ["парсер"])
    assert not matches(logo_order, ["парсер", "python"])


def test_keyword_matches_word_start_only():
    order = {"title": "Разработать проект котельной", "description": "", "category": ""}
    assert not matches(order, ["бот"])
    order["title"] = "Нужны боты для Telegram"
    assert matches(order, ["бот"])


def test_extract_budget():
    assert extract_budget("Бюджет: 15000 руб") == 15000
    assert extract_budget("Цена договорная") is None


def test_check_collects_all_feeds_and_survives_failed_one(tmp_path, monkeypatch):
    import monitor
    import requests

    class Response:
        content = SAMPLE.encode()

        def raise_for_status(self):
            pass

    def fake_get(url, **kwargs):
        if "habr" in url:
            raise requests.ConnectionError("нет сети")
        return Response()

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(monitor, "SEEN_FILE", tmp_path / "seen.json")
    monkeypatch.setattr(monitor, "XLSX_FILE", tmp_path / "orders.xlsx")
    monkeypatch.setattr(monitor.requests, "get", fake_get)
    monkeypatch.setattr(monitor, "FEEDS", {"FL.ru": "https://fl.example/rss",
                                           "Хабр": "https://habr.example/rss"})

    monitor.check(["парсер"], telegram=False)
    from openpyxl import load_workbook
    rows = list(load_workbook(tmp_path / "orders.xlsx").active.values)
    assert rows[0][-1] == "source"
    assert len(rows) == 2 and rows[1][-1] == "FL.ru"

    monitor.check(["парсер"], telegram=False)  # повторный запуск не дублирует
    assert len(list(load_workbook(tmp_path / "orders.xlsx").active.values)) == 2


def test_telegram_error_hides_token_and_keeps_undelivered(tmp_path, monkeypatch):
    import monitor
    import requests

    two_items = SAMPLE.replace("Нарисовать логотип", "Нужен парсер логотипов")

    class Response:
        content = two_items.encode()

        def raise_for_status(self):
            pass

    sent = []

    def fake_post(url, data, **kwargs):
        if sent:  # первое сообщение уходит, на втором связь рвётся
            raise requests.ConnectionError(f"обрыв при запросе к {url}")
        sent.append(data["text"])
        return Response()

    monkeypatch.setenv("BOT_TOKEN", "123:SECRET")
    monkeypatch.setenv("ADMIN_ID", "42")
    monkeypatch.setattr(monitor, "SEEN_FILE", tmp_path / "seen.json")
    monkeypatch.setattr(monitor, "XLSX_FILE", tmp_path / "orders.xlsx")
    monkeypatch.setattr(monitor, "FEEDS", {"FL.ru": "https://fl.example/rss"})
    monkeypatch.setattr(monitor.requests, "get", lambda *a, **k: Response())
    monkeypatch.setattr(monitor.requests, "post", fake_post)
    monkeypatch.setattr(monitor.time, "sleep", lambda s: None)

    monitor.check(["парсер"], telegram=True)
    assert len(sent) == 1
    assert len(monitor.load_seen()) == 1  # недоставленный заказ не помечен как виденный

    import io, contextlib
    out = io.StringIO()
    sent.append("уже был")
    with contextlib.redirect_stdout(out):
        monitor.check(["парсер"], telegram=True)
    assert "SECRET" not in out.getvalue() and "***" in out.getvalue()
