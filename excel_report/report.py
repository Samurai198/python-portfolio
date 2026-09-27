"""Автоматический отчёт по продажам из CSV/Excel.

Берёт таблицу продаж (дата, менеджер, товар, количество, цена),
считает выручку и собирает Excel-отчёт из нескольких листов:
итоги по менеджерам, по товарам, по месяцам и исходные данные.

Запуск:
    python report.py sales.csv --out report.xlsx
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

REQUIRED = ["date", "manager", "product", "qty", "price"]
HEADERS = {"date": "Дата", "manager": "Менеджер", "product": "Товар", "qty": "Кол-во",
           "price": "Цена", "revenue": "Выручка", "deals": "Сделок", "month": "Месяц"}


def load(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xls"):
        df = pd.read_excel(path)
    else:
        df = pd.read_csv(path, sep=None, engine="python")
    df.columns = [c.strip().lower() for c in df.columns]
    missing = set(REQUIRED) - set(df.columns)
    if missing:
        raise ValueError(f"В файле не хватает колонок: {', '.join(sorted(missing))}")
    return df


def build(df):
    df = df.dropna(subset=["qty", "price"]).copy()
    df["date"] = pd.to_datetime(df["date"], dayfirst=True)
    df["revenue"] = df["qty"] * df["price"]
    df["month"] = df["date"].dt.to_period("M").astype(str)

    by_manager = (df.groupby("manager")
                    .agg(deals=("revenue", "size"), revenue=("revenue", "sum"))
                    .sort_values("revenue", ascending=False)
                    .reset_index())
    by_product = (df.groupby("product")
                    .agg(qty=("qty", "sum"), revenue=("revenue", "sum"))
                    .sort_values("revenue", ascending=False)
                    .reset_index())
    by_month = df.groupby("month")["revenue"].sum().reset_index()
    return {"Менеджеры": by_manager, "Товары": by_product,
            "По месяцам": by_month, "Данные": df.drop(columns="month")}


def save(sheets, out):
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, table in sheets.items():
            table.rename(columns=HEADERS).to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if hasattr(cell.value, "year"):
                        cell.number_format = "DD.MM.YYYY"
                    elif isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0"
            for column in ws.columns:
                width = max(len(str(c.value or "")) for c in column) + 2
                ws.column_dimensions[column[0].column_letter].width = min(width, 50)


def main():
    ap = argparse.ArgumentParser(description="Отчёт по продажам в Excel")
    ap.add_argument("input", help="CSV или XLSX с продажами")
    ap.add_argument("--out", default="report.xlsx")
    args = ap.parse_args()

    sheets = build(load(args.input))
    save(sheets, args.out)
    total = sheets["Менеджеры"]["revenue"].sum()
    print(f"Отчёт сохранён в {args.out}. Общая выручка: {total:,.0f}".replace(",", " "))


if __name__ == "__main__":
    main()
